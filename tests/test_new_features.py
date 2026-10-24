import unittest

try:
    import docx as _docx  # noqa: F401

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pypdf  # noqa: F401

    HAS_PDF = True
except ImportError:
    HAS_PDF = False

from chunkcite.chunking import Chunk, audit, chunk_file, chunk_text


def _make_pdf(path, texts):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(612, 792)
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
        font = DictionaryObject()
        font[NameObject("/Type")] = NameObject("/Font")
        font[NameObject("/Subtype")] = NameObject("/Type1")
        font[NameObject("/BaseFont")] = NameObject("/Helvetica")
        fonts = DictionaryObject()
        fonts[NameObject("/F1")] = writer._add_object(font)
        resources = DictionaryObject()
        resources[NameObject("/Font")] = fonts
        page[NameObject("/Resources")] = resources
    with open(path, "wb") as f:
        writer.write(f)


class TestAuditPages(unittest.TestCase):
    def test_audit_surfaces_pages(self):
        chunks = [
            Chunk(id="a" * 12, text="one", source="d.pdf", heading_path=["H"],
                  start_line=1, end_line=1, page=1),
            Chunk(id="b" * 12, text="two", source="d.pdf", heading_path=["H"],
                  start_line=2, end_line=2, page=2),
            Chunk(id="c" * 12, text="three", source="d.md", heading_path=["H"],
                  start_line=1, end_line=1, page=None),
        ]
        report = audit(chunks)
        self.assertEqual(report["pages"], [1, 2])
        self.assertEqual(report["chunks_by_page"], {"1": 1, "2": 1})
        self.assertEqual(report["chunks_without_page"], 1)

    @unittest.skipUnless(HAS_PDF, "pypdf not installed")
    def test_pdf_chunks_carry_page_numbers(self):
        import tempfile
        from pathlib import Path

        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "doc.pdf"
            _make_pdf(pdf, ["Hello page one pgvector", "Hello page two citations"])
            chunks = chunk_file(pdf)
            self.assertEqual([c.page for c in chunks], [1, 2])
            report = audit(chunks)
            self.assertEqual(report["pages"], [1, 2])


@unittest.skipUnless(HAS_DOCX, "python-docx not installed")
class TestDocx(unittest.TestCase):
    def test_docx_headings_and_tables(self):
        import tempfile
        from pathlib import Path

        import docx

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "doc.docx"
            document = docx.Document()
            document.add_heading("Retrieval", level=1)
            document.add_paragraph("pgvector adds vector search to Postgres.")
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "retriever"
            table.rows[0].cells[1].text = "strength"
            table.rows[1].cells[0].text = "BM25"
            table.rows[1].cells[1].text = "exact terms"
            document.save(path)
            chunks = chunk_file(path)
            self.assertTrue(any("pgvector" in c.text for c in chunks))
            tables = [c for c in chunks if c.kind == "table"]
            self.assertEqual(len(tables), 1)
            self.assertIn("BM25", tables[0].text)
            self.assertEqual(tables[0].heading_path, ["Retrieval"])


class TestGeminiAdapter(unittest.TestCase):
    def test_answer_and_verify_flags_invented_claim(self):
        from chunkcite.adapters.gemini import answer_and_verify

        chunks = chunk_text(
            "# Facts\n\npgvector adds vector similarity search to Postgres.\n",
            "doc.md",
        )
        cid = chunks[0].id

        def fake_generate(prompt):
            self.assertIn(cid, prompt)
            return (
                f"pgvector adds vector similarity search to Postgres [{cid}]. "
                f"pgvector is 99 percent faster than everything [{cid}]."
            )

        result = answer_and_verify(
            "What does pgvector add?", chunks, fake_generate
        )
        self.assertEqual(result["total"], 2)
        self.assertEqual(result["supported"], 1)
        verdicts = [v["verdict"] for v in result["verification"]]
        self.assertEqual(verdicts, ["supported", "unsupported"])

    def test_uncited_sentence_is_unresolved(self):
        from chunkcite.adapters.gemini import parse_claims, answer_and_verify
        from chunkcite.chunking import chunk_text

        self.assertEqual(
            parse_claims("A claim with no citation."),
            [{"text": "A claim with no citation.", "cite": None}],
        )
        chunks = chunk_text("# F\n\nReal text here.\n", "doc.md")
        result = answer_and_verify("Q?", chunks, lambda prompt: "No citation here.")
        self.assertEqual(result["verification"][0]["verdict"], "unresolved")


if __name__ == "__main__":
    unittest.main()
